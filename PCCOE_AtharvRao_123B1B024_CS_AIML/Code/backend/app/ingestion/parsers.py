"""Deterministic parsers; parser failures are isolated to the affected unit where possible."""

from __future__ import annotations

import csv
import io
import json
import re
from abc import ABC, abstractmethod

import fitz
from docx import Document as DocxDocument
from openpyxl import load_workbook

from app.ingestion.contracts import (
    DocumentFormat,
    IngestionError,
    ParsedBlock,
    ParsedDocument,
    SourcePosition,
    TableData,
)


class OCRProvider(ABC):
    """Optional local OCR boundary for scanned pages; no cloud fallback is permitted."""

    @abstractmethod
    def extract_text(self, image_bytes: bytes, page_number: int) -> str:
        """Extract text from a rendered document page."""


class DocumentParser(ABC):
    @abstractmethod
    def parse(self, content: bytes) -> ParsedDocument:
        """Parse raw source bytes into page/position-aware blocks."""


class PDFParser(DocumentParser):
    def __init__(self, ocr_provider: OCRProvider | None = None) -> None:
        self._ocr_provider = ocr_provider

    def parse(self, content: bytes) -> ParsedDocument:
        try:
            pdf = fitz.open(stream=content, filetype="pdf")
        except Exception as error:
            raise IngestionError("PDF parser could not open uploaded source.") from error
        blocks: list[ParsedBlock] = []
        warnings: list[str] = []
        for page_index, page in enumerate(pdf, start=1):
            try:
                text = page.get_text("text").strip()
                if not text and self._ocr_provider is not None:
                    pixmap = page.get_pixmap(dpi=150, alpha=False)
                    text = self._ocr_provider.extract_text(
                        pixmap.tobytes("png"), page_index
                    ).strip()
                if not text:
                    warnings.append(f"Page {page_index} contains no extractable text.")
                tables = self._tables_for_page(page, page_index)
                blocks.append(
                    ParsedBlock(text=text, position=SourcePosition(page=page_index), tables=tables)
                )
            except Exception as error:
                warnings.append(
                    f"Page {page_index} could not be fully parsed: {type(error).__name__}."
                )
                blocks.append(
                    ParsedBlock(
                        text="", position=SourcePosition(page=page_index), warning=str(error)
                    )
                )
        return ParsedDocument(blocks=blocks, warnings=warnings)

    @staticmethod
    def _tables_for_page(page: fitz.Page, page_number: int) -> list[TableData]:
        try:
            found = page.find_tables()
            return [
                TableData(
                    rows=[[str(cell or "") for cell in row] for row in table.extract()],
                    position=SourcePosition(page=page_number),
                )
                for table in found.tables
            ]
        except Exception:
            return []


class DOCXParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            document = DocxDocument(io.BytesIO(content))
        except Exception as error:
            raise IngestionError("DOCX parser could not open uploaded source.") from error
        blocks = [
            ParsedBlock(text=paragraph.text, position=SourcePosition())
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            blocks.append(
                ParsedBlock(
                    text="", position=SourcePosition(), tables=[TableData(rows, SourcePosition())]
                )
            )
        return ParsedDocument(blocks=blocks)


class XLSXParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise IngestionError("XLSX parser could not open uploaded source.") from error
        blocks: list[ParsedBlock] = []
        for worksheet in workbook.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            blocks.append(
                ParsedBlock(
                    text=f"# {worksheet.title}",
                    position=SourcePosition(),
                    tables=[TableData(rows, SourcePosition())],
                )
            )
        return ParsedDocument(blocks=blocks)


class CSVParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            decoded = content.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(decoded)))
        except (UnicodeDecodeError, csv.Error) as error:
            raise IngestionError("CSV parser could not read uploaded source.") from error
        return ParsedDocument(
            blocks=[
                ParsedBlock(
                    text="", position=SourcePosition(), tables=[TableData(rows, SourcePosition())]
                )
            ]
        )


class MarkdownParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as error:
            raise IngestionError("Markdown source must be UTF-8.") from error
        return ParsedDocument(
            blocks=[
                ParsedBlock(
                    text=text, position=SourcePosition(start_offset=0, end_offset=len(text))
                )
            ]
        )


class ParserFactory:
    """Maps validated document formats to deterministic parser implementations."""

    def __init__(self, ocr_provider: OCRProvider | None = None) -> None:
        self._parsers: dict[DocumentFormat, DocumentParser] = {
            DocumentFormat.PDF: PDFParser(ocr_provider),
            DocumentFormat.DOCX: DOCXParser(),
            DocumentFormat.XLSX: XLSXParser(),
            DocumentFormat.CSV: CSVParser(),
            DocumentFormat.MARKDOWN: MarkdownParser(),
        }

    def parser_for(self, document_format: DocumentFormat) -> DocumentParser:
        return self._parsers[document_format]


def table_as_evidence(table: TableData) -> str:
    """Preserve parsed tables in chunk content as explicitly labeled structured JSON evidence."""
    return "[TABLE] " + json.dumps(table.rows, ensure_ascii=False)


# Numbered engineering headings start with a positive section number.  Requiring
# positive numeric segments prevents table values such as "0.00 to 320.00" from
# being promoted to headings during PDF normalization.
HEADING_PATTERN = re.compile(
    r"^(?P<number>[1-9]\d*(?:\.[1-9]\d*)*\.?)[\s]+(?P<title>.+)$"
    r"|^(?P<hash>#{1,6})\s+(?P<markdown>.+)$"
)
