"""Builders for realistic test documents (PDF with layout, Word with styles, scanned pages)."""

from __future__ import annotations

import io

import fitz  # PyMuPDF
from docx import Document as DocxDocument

BODY = 10
H1 = 18
H2 = 14
HEADER_TEXT = "AUTOMOTIVE ENGINEERING AI | PROJECT CASE STUDIES"


def _footer(page: fitz.Page, number: int) -> None:
    page.insert_text(
        (72, 815),
        f"Confidential | Reference Solution Document Page {number}",
        fontsize=8,
    )


def _header(page: fitz.Page) -> None:
    page.insert_text((72, 40), HEADER_TEXT, fontsize=8)


def _paragraph(page: fitz.Page, top: float, text: str, height: float = 60) -> float:
    page.insert_textbox(
        fitz.Rect(72, top, 523, top + height), text, fontsize=BODY, fontname="helv"
    )
    return top + height + 6


def make_nested_pdf() -> bytes:
    """Three pages: running header/footer, sized headings, wrapped prose, tricky numeric lines."""
    document = fitz.open()
    # page 1
    page = document.new_page()
    _header(page)
    page.insert_text((72, 100), "1. Brake Subsystem", fontsize=H1, fontname="hebo")
    top = _paragraph(
        page,
        110,
        "The controller reads wheel speed from the sensor over 3 channels that are sampled "
        "every 10 ms. The supply is 12 V nominal. Values are validated before use.",
        70,
    )
    page.insert_text(
        (72, top + 14), "1.1 Wheel Speed Interface", fontsize=H2, fontname="hebo"
    )
    _paragraph(
        page,
        top + 20,
        "The interface carries the WheelSpeedSignal to the controller.",
        40,
    )
    _footer(page, 1)
    # page 2
    page = document.new_page()
    _header(page)
    page.insert_text(
        (72, 100), "1.2 Brake Command Interface", fontsize=H2, fontname="hebo"
    )
    _paragraph(page, 110, "The BrakeController provides the BrakeCommandInterface.", 40)
    page.insert_text((72, 200), "2. Diagnostics", fontsize=H1, fontname="hebo")
    _paragraph(
        page, 210, "Diagnostic services follow ISO 14229 and use session 0x03.", 40
    )
    _footer(page, 2)
    # page 3
    page = document.new_page()
    _header(page)
    _paragraph(page, 100, "Faults are stored in non-volatile memory.", 40)
    _footer(page, 3)
    return document.tobytes()


def make_toc_pdf() -> bytes:
    """A PDF whose bookmarks define the hierarchy; body headings share the body font size."""
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 100), "Introduction", fontsize=BODY, fontname="helv")
    _paragraph(page, 110, "This document describes the brake controller.", 30)
    page.insert_text((72, 200), "Scope", fontsize=BODY, fontname="helv")
    _paragraph(page, 210, "Only the software architecture is covered.", 30)
    page = document.new_page()
    page.insert_text((72, 100), "Interfaces", fontsize=BODY, fontname="helv")
    _paragraph(page, 110, "The controller uses the WheelSpeedInterface.", 30)
    document.set_toc([[1, "Introduction", 1], [2, "Scope", 1], [1, "Interfaces", 2]])
    return document.tobytes()


def make_scanned_pdf(
    text: str = "BrakeController requires WheelSpeedInterface",
) -> bytes:
    """A page that contains only a raster image (no text layer), like a scan."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (800, 200), "white")
    ImageDraw.Draw(image).text((20, 80), text, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    document = fitz.open()
    page = document.new_page()
    page.insert_image(fitz.Rect(50, 50, 550, 175), stream=buffer.getvalue())
    return document.tobytes()


def make_toc_pdf_with_contents_page() -> bytes:
    """Page 1 is a printed table of contents with dot leaders; real content follows."""
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 100), "Contents", fontsize=H1, fontname="hebo")
    page.insert_text((72, 140), "1. Overview ........................ 2", fontsize=BODY)
    page.insert_text((72, 160), "2. Interfaces ...................... 3", fontsize=BODY)
    page = document.new_page()
    page.insert_text((72, 100), "1. Overview", fontsize=H1, fontname="hebo")
    _paragraph(page, 110, "Overview text about the BrakeController.", 30)
    return document.tobytes()


def make_docx_with_styles() -> bytes:
    document = DocxDocument()
    document.add_heading("Brake Controller HLD", level=0)  # style "Title"
    document.add_paragraph("Document Version: 2.1.0")
    document.add_heading("Architecture", level=1)
    document.add_paragraph("The BrakeController reads the WheelSpeedSensor.")
    document.add_heading("Components", level=2)
    document.add_paragraph("WheelSpeedSensor", style="List Bullet")
    document.add_paragraph("ABS_Module", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Signal"
    table.rows[0].cells[1].text = "Type"
    table.rows[1].cells[0].text = "Speed"
    table.rows[1].cells[1].text = "uint16"
    document.add_paragraph("Text after the table.")
    document.add_heading("Interfaces", level=1)
    document.add_paragraph("The WheelSpeedInterface carries speed.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
