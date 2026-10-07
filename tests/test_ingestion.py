"""Deterministic ingestion/parser coverage without a running PostgreSQL instance."""

from __future__ import annotations

import io

import fitz
import pytest
from app.ingestion.chunking import chunk_section
from app.ingestion.contracts import (
    IngestionError,
    NormalizedSection,
    ParsedBlock,
    SourcePosition,
)
from app.ingestion.normalize import normalize_document
from app.ingestion.parsers import OCRProvider, ParserFactory, PDFParser
from app.ingestion.validation import content_hash, safe_filename, validate_upload
from docx import Document as DocxDocument
from openpyxl import Workbook


def make_pdf(lines: list[str]) -> bytes:
    document = fitz.open()
    page = document.new_page()
    for index, line in enumerate(lines):
        page.insert_text((72, 72 + index * 20), line)
    return document.tobytes()


def test_normal_pdf_preserves_page_and_text() -> None:
    parsed = PDFParser().parse(
        make_pdf(["1 Architecture", "BodyControlComponent provides status."])
    )

    assert parsed.blocks[0].position.page == 1
    assert "BodyControlComponent" in parsed.blocks[0].text


def test_sectioned_pdf_preserves_hierarchy() -> None:
    parsed = PDFParser().parse(
        make_pdf(["1 System", "1.1 Components", "BrakeController details"])
    )
    sections = normalize_document(parsed)

    system = next(section for section in sections if section.heading == "System")
    assert system.children[0].heading == "Components"
    assert system.children[0].content[0].text == "BrakeController details"


def test_table_measurements_are_not_interpreted_as_numbered_headings() -> None:
    parsed = PDFParser().parse(
        make_pdf(["6 Data Elements", "0.00 to 320.00", "WheelSpeed"])
    )

    sections = normalize_document(parsed)
    data_elements = next(
        section for section in sections if section.heading == "Data Elements"
    )

    assert data_elements.children == []
    assert [block.text for block in data_elements.content] == [
        "0.00 to 320.00",
        "WheelSpeed",
    ]


def test_table_pdf_preserves_structured_rows() -> None:
    document = fitz.open()
    page = document.new_page()
    page.draw_rect(fitz.Rect(72, 72, 272, 132))
    page.draw_line((172, 72), (172, 132))
    page.draw_line((72, 102), (272, 102))
    for point, text in [
        ((80, 90), "Signal"),
        ((180, 90), "Type"),
        ((80, 120), "Speed"),
        ((180, 120), "uint16"),
    ]:
        page.insert_text(point, text)
    parsed = PDFParser().parse(document.tobytes())

    assert parsed.blocks[0].tables[0].rows == [["Signal", "Type"], ["Speed", "uint16"]]


class FakeOCR(OCRProvider):
    def extract_text(self, image_bytes: bytes, page_number: int) -> str:
        assert image_bytes
        return f"OCR page {page_number}"


def test_scanned_pdf_uses_injected_local_ocr_fallback() -> None:
    blank = fitz.open()
    blank.new_page()
    parsed = PDFParser(FakeOCR()).parse(blank.tobytes())

    assert parsed.blocks[0].text == "OCR page 1"


def test_malformed_input_and_mismatched_mime_are_rejected() -> None:
    with pytest.raises(IngestionError):
        validate_upload("design.pdf", "application/pdf", b"not-a-pdf", 100)
    with pytest.raises(IngestionError):
        validate_upload("design.pdf", "text/plain", b"%PDF-1.7", 100)


def test_filename_is_never_trusted_as_a_path_and_hash_is_stable() -> None:
    assert safe_filename("../../confidential/HLD.pdf") == "HLD.pdf"
    assert content_hash(b"source") == content_hash(b"source")


def test_docx_xlsx_csv_and_markdown_parsers() -> None:
    docx = DocxDocument()
    docx.add_heading("System", level=1)
    docx.add_paragraph("Document content")
    docx_buffer = io.BytesIO()
    docx.save(docx_buffer)
    assert (
        "Document content"
        in ParserFactory()
        .parser_for(validate_upload("hld.docx", None, docx_buffer.getvalue(), 100000))
        .parse(docx_buffer.getvalue())
        .blocks[1]
        .text
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Signals"
    worksheet.append(["Signal", "Type"])
    worksheet.append(["Speed", "uint16"])
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)
    xlsx = xlsx_buffer.getvalue()
    assert (
        ParserFactory()
        .parser_for(validate_upload("catalog.xlsx", None, xlsx, 100000))
        .parse(xlsx)
        .blocks[0]
        .tables[0]
        .rows[1][0]
        == "Speed"
    )

    csv_source = b"Signal,Type\nSpeed,uint16\n"
    assert ParserFactory().parser_for(
        validate_upload("signals.csv", "text/csv", csv_source, 100000)
    ).parse(csv_source).blocks[0].tables[0].rows[0] == ["Signal", "Type"]
    markdown = b"# Architecture\nBrakeController\n"
    assert (
        ParserFactory()
        .parser_for(validate_upload("hld.md", "text/markdown", markdown, 100000))
        .parse(markdown)
        .blocks[0]
        .text.startswith("# Architecture")
    )


def test_hierarchy_preserving_chunking_only_splits_oversized_sections() -> None:
    section = NormalizedSection(
        heading="Interfaces",
        level=2,
        position=SourcePosition(page=4),
        content=[
            ParsedBlock("one two three four five six seven", SourcePosition(page=4))
        ],
    )
    chunks = chunk_section(section, max_tokens=4, overlap_tokens=1)

    assert len(chunks) == 2
    assert all(chunk.content.startswith("Section: Interfaces") for chunk in chunks)
    assert all(chunk.position.page == 4 for chunk in chunks)
