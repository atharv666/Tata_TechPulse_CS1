"""Structured export and the entity report, over HTTP."""

from __future__ import annotations

import csv
import io

import pytest
from app.main import app
from app.services.export import safe_cell
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.support.fakes import installed_fakes
from tests.support.pipeline import drain, real_runner

pytestmark = pytest.mark.integration

ALICE = {"X-User-Id": "alice"}  # project ADMIN
BOB = {"X-User-Id": "bob"}  # VIEWER
CAROL = {"X-User-Id": "carol"}  # REVIEWER

DOCUMENT = b"""# Brake HLD
Document Version: 1.0.0
Status: Approved

## 1 Components
The primary software component is `BrakeController` (also referenced as `BrakeCtrl`).

- `BrakeController` is a `SOFTWARE_COMPONENT` that manages braking.
- `BrakeCtrl` is a `SOFTWARE_COMPONENT` that appears in the wiring diagram.
- `WheelSpeedSensor` is a `SOFTWARE_COMPONENT` that reads wheel speed.
- `WheelSpeedInterface` is an `INTERFACE` that carries wheel speed data.
- `=SUM(A1:A2)` is a `SIGNAL` that is named like a spreadsheet formula.

## 2 Dependencies
- `BrakeController` `REQUIRES` `WheelSpeedInterface` through input port `P_Wheel_In`.
- `WheelSpeedSensor` `PROVIDES` `WheelSpeedInterface`.
- `BrakeController` `PART_OF` `ChassisDomain`.
"""
QUESTION = "What does BrakeController require?"


@pytest.fixture
def system(clean_database):
    with TestClient(app, raise_server_exceptions=False) as client:
        project_id = client.post(
            "/api/v1/projects", json={"name": "admin-view"}, headers=ALICE
        ).json()["id"]
        base = f"/api/v1/projects/{project_id}"
        for user, role in (("bob", "VIEWER"), ("carol", "REVIEWER")):
            client.post(
                f"{base}/members", json={"user_id": user, "role": role}, headers=ALICE
            )
        client.post(
            f"{base}/documents/upload",
            headers=ALICE,
            files={"file": ("Brake_HLD.md", DOCUMENT, "text/markdown")},
        )
        with installed_fakes():
            drain(real_runner())
            for headers, mode in ((ALICE, "HYBRID"), (BOB, "GRAPH"), (BOB, "SOURCE")):
                answer = client.post(
                    f"{base}/query",
                    headers=headers,
                    json={"query": QUESTION, "mode": mode},
                )
                assert answer.status_code == 200, answer.text
            yield client, base


def get(
    client: TestClient, url: str, headers: dict[str, str] = ALICE, **params: object
):
    response = client.get(url, headers=headers, params=params)
    assert response.status_code == 200, response.text
    return response.json()


# -- export --------------------------------------------------------------------------------------
def test_json_export_has_provenance_and_the_live_graph_only(system) -> None:
    client, base = system
    data = get(client, f"{base}/export", format="json")

    assert data["counts"]["entities"] >= 4 and "generated_at" in data
    names = {e["name"] for e in data["entities"]}
    assert "BrakeController" in names and "BrakeCtrl" not in names  # merged away
    controller = next(e for e in data["entities"] if e["name"] == "BrakeController")
    assert controller["aliases"] == ["BrakeCtrl"]
    assert (
        controller["source_document"] == "Brake_HLD.md"
        and controller["evidence_count"] >= 1
    )
    assert controller["confidence"] > 0 and controller["extraction_type"] == "EXPLICIT"
    assert controller["validation_state"] == "PENDING"
    edge = next(
        r for r in data["relationships"] if r["relationship_type"] == "REQUIRES"
    )
    assert (edge["source"], edge["target"]) == (
        "BrakeController",
        "WheelSpeedInterface",
    )
    assert "REQUIRES" in edge["source_excerpt"] and edge["source_section"].endswith(
        "2 Dependencies"
    )
    assert data["documents"][0]["declared_status"] == "Approved"
    component = next(
        c for c in data["components"] if c["component"] == "BrakeController"
    )
    assert component["requires"] == ["WheelSpeedInterface"]
    kinds = {f["finding_type"] for f in data["findings"]}
    assert "MISSING_ENTITY_DEFINITION" in kinds  # ChassisDomain is never defined


def test_csv_export_is_flat_and_safe_to_open_in_a_spreadsheet(system) -> None:
    client, base = system
    response = client.get(
        f"{base}/export", headers=BOB, params={"format": "csv", "table": "entities"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "entities.csv" in response.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert {"name", "type", "confidence", "validation_state", "source_document"} <= set(
        rows[0]
    )
    formula = next(r for r in rows if "SUM" in r["name"])
    assert formula["name"].startswith("'=")  # a leading '=' is neutralised
    assert (
        safe_cell("=1+1") == "'=1+1"
        and safe_cell("plain") == "plain"
        and safe_cell(3) == 3
    )
    aliases = next(r for r in rows if r["name"] == "BrakeController")["aliases"]
    assert aliases == "BrakeCtrl"


def test_xlsx_export_has_a_sheet_per_table(system) -> None:
    client, base = system
    response = client.get(f"{base}/export", headers=ALICE, params={"format": "xlsx"})
    assert (
        response.status_code == 200
        and "spreadsheetml" in response.headers["content-type"]
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [
        "Entities",
        "Relationships",
        "Findings",
        "Components",
    ]
    header = [cell.value for cell in workbook["Relationships"][1]]
    assert {
        "source",
        "relationship_type",
        "target",
        "confidence",
        "source_excerpt",
    } <= set(header)
    assert workbook["Entities"].max_row >= 5


def test_rejected_facts_are_excluded_unless_asked_for(system) -> None:
    client, base = system
    sensor = get(client, f"{base}/entities", q="WheelSpeedSensor")["items"][0]
    client.post(
        f"{base}/entities/{sensor['id']}/reject",
        headers=CAROL,
        json={"reason": "not real"},
    )

    default = {
        e["name"] for e in get(client, f"{base}/export", format="json")["entities"]
    }
    included = {
        e["name"]
        for e in get(client, f"{base}/export", format="json", include_rejected="true")[
            "entities"
        ]
    }
    assert "WheelSpeedSensor" not in default and "WheelSpeedSensor" in included


def test_export_parameters_are_validated(system) -> None:
    client, base = system
    assert (
        client.get(
            f"{base}/export", headers=ALICE, params={"format": "pdf"}
        ).status_code
        == 422
    )
    assert (
        client.get(
            f"{base}/export",
            headers=ALICE,
            params={"format": "csv", "table": "secrets"},
        ).status_code
        == 422
    )


# -- entity report -----------------------------------------------------------------------------------
def test_the_entity_report_groups_what_the_graph_states_with_sources(system) -> None:
    client, base = system
    controller = get(client, f"{base}/entities", q="BrakeController")["items"][0]

    report = get(client, f"{base}/entities/{controller['id']}/report")

    assert report["entity"]["canonical_name"] == "BrakeController"
    assert report["aliases"] == ["BrakeCtrl"]
    [requires] = report["outgoing"]["REQUIRES"]
    assert (
        requires["name"] == "WheelSpeedInterface"
        and requires["entity_type"] == "INTERFACE"
    )
    assert (
        requires["source_document"] == "Brake_HLD.md"
        and "REQUIRES" in requires["source_excerpt"]
    )
    assert requires["status_text"].endswith("EXPLICIT · PENDING")
    interface = get(client, f"{base}/entities", q="WheelSpeedInterface")["items"][0]
    incoming = get(client, f"{base}/entities/{interface['id']}/report")["incoming"]
    assert {i["name"] for i in incoming["REQUIRES"]} == {"BrakeController"}
    assert {i["name"] for i in incoming["PROVIDES"]} == {"WheelSpeedSensor"}
    missing = "00000000-0000-0000-0000-000000000000"
    assert (
        client.get(f"{base}/entities/{missing}/report", headers=ALICE).status_code
        == 404
    )
