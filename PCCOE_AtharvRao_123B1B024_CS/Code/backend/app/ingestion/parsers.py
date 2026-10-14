"""Deterministic, layout-aware parsers. No model is involved at any point.

Each parser returns ``ParsedBlock`` objects in reading order: headings (when the format itself marks
them), paragraphs and tables. What makes a document *nested* is decided in this order of trust:

1. PDF bookmarks / Word heading styles / Markdown ``#`` (the author said so),
2. PDF font size and weight (the layout says so),
3. numbering such as ``1.2.3`` (only as a validated fallback, see ``normalize``).

Page numbers, running headers/footers and dot-leader table-of-contents lines are removed from PDFs
so they cannot leak into chunks, embeddings or the lexical index.
"""

from __future__ import annotations

import csv
import io
import re
from abc import ABC, abstractmethod
from collections import Counter
from dataclasses import dataclass, field
from typing import Any

try:  # PyMuPDF >= 1.24 exposes ``pymupdf``; ``fitz`` is the legacy alias
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz  # type: ignore[no-redef]

from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook

from app.ingestion.contracts import (
    BlockKind,
    DocumentFormat,
    IngestionError,
    ParsedBlock,
    ParsedDocument,
    SourcePosition,
    TableData,
)
from app.ingestion.structure import looks_like_heading_text, number_depth, split_number
from app.ingestion.text import LIST_MARKER, is_bullet, is_list_item, render_table_markdown


@dataclass(frozen=True)
class OCRResult:
    text: str
    confidence: float | None = None
    warnings: tuple[str, ...] = ()


class OCRProvider(ABC):
    """Boundary for turning a rendered page image into text.

    ``external`` is True when page images leave this machine (a hosted vision API). The ingestion
    service refuses external providers unless the deployment explicitly allows them.
    """

    name = "ocr"
    external = False

    @abstractmethod
    def extract_text(self, image_bytes: bytes, page_number: int) -> str:
        """Extract text from a rendered document page (PNG bytes)."""

    def extract(self, image_bytes: bytes, page_number: int) -> OCRResult:
        return OCRResult(self.extract_text(image_bytes, page_number))


class DocumentParser(ABC):
    @abstractmethod
    def parse(self, content: bytes) -> ParsedDocument:
        """Parse raw source bytes into page/position-aware blocks."""


# --------------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------------
_PAGE_NUMBER_LINE = re.compile(
    r"^(?:page\s+)?\d{1,4}(?:\s*(?:/|of)\s*\d{1,4})?$|^[-–—]\s*\d{1,4}\s*[-–—]$",
    re.IGNORECASE,
)
_DOT_LEADER = re.compile(r"(?:\.\s?){4,}\s*\d{1,4}\s*$")
_BOLD_FLAG = 16
_EDGE_ZONE = 0.09  # top/bottom share of the page height where running lines live


@dataclass
class _Line:
    text: str
    size: float
    bold: bool
    x0: float
    y0: float
    y1: float
    block: int
    page: int
    drop: bool = False
    heading_level: int | None = None
    heading_source: str | None = None


@dataclass
class _PageData:
    number: int
    height: float
    lines: list[_Line] = field(default_factory=list)
    tables: list[tuple[float, list[list[str]], str | None]] = field(default_factory=list)
    has_images: bool = False


def _normalise_key(text: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", text.strip().lower()))


def _title_key(text: str) -> str:
    _, title = split_number(text)
    return re.sub(r"[^a-z0-9]+", "", title.lower())


class PDFParser(DocumentParser):
    def __init__(
        self,
        ocr_provider: OCRProvider | None = None,
        ocr_min_chars: int = 30,
        ocr_dpi: int = 200,
        ocr_max_pages: int | None = None,
    ) -> None:
        self._ocr_provider = ocr_provider
        self._ocr_min_chars = ocr_min_chars
        self._ocr_dpi = ocr_dpi
        self._ocr_max_pages = ocr_max_pages

    # -- public ------------------------------------------------------------------------------
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            pdf = fitz.open(stream=content, filetype="pdf")  # type: ignore[no-untyped-call]
        except Exception as error:
            raise IngestionError("PDF parser could not open uploaded source.") from error
        warnings: list[str] = []
        pages = self._collect_pages(pdf, warnings)
        removed = self._drop_running_lines(pages)
        toc = self._table_of_contents(pdf)
        structure = self._mark_headings(pages, toc)
        blocks, ocr_pages = self._build_blocks(pdf, pages, warnings)
        metadata = self._metadata(pdf, len(pages))
        return ParsedDocument(
            blocks=blocks,
            warnings=warnings,
            structure_source=structure,
            metadata=metadata,
            ocr_pages=ocr_pages,
            removed_running_lines=removed,
        )

    # -- extraction --------------------------------------------------------------------------
    def _collect_pages(self, pdf: Any, warnings: list[str]) -> list[_PageData]:
        pages: list[_PageData] = []
        for number, page in enumerate(pdf, start=1):
            data = _PageData(number=number, height=float(page.rect.height))
            table_boxes: list[Any] = []
            try:
                for table in page.find_tables().tables:
                    rows = [[str(cell or "") for cell in row] for row in table.extract()]
                    if rows:
                        table_boxes.append(table.bbox)
                        data.tables.append((float(table.bbox[1]), rows, None))
            except Exception as error:  # table detection is best effort; text is still parsed
                warnings.append(
                    f"Tables on page {number} could not be detected: {type(error).__name__}."
                )
            try:
                text_dict = page.get_text("dict", sort=True)
                data.has_images = any(b.get("type") == 1 for b in text_dict["blocks"])
                for block in text_dict["blocks"]:
                    if block.get("type") != 0:
                        continue
                    for line in block["lines"]:
                        spans = [s for s in line["spans"] if s["text"].strip()]
                        if not spans:
                            continue
                        x0, y0, x1, y1 = line["bbox"]
                        if any(
                            b[0] - 2 <= (x0 + x1) / 2 <= b[2] + 2
                            and b[1] - 2 <= (y0 + y1) / 2 <= b[3] + 2
                            for b in table_boxes
                        ):
                            continue  # text that belongs to a table is emitted as that table
                        size = max(float(s["size"]) for s in spans)
                        bold = all(
                            (int(s["flags"]) & _BOLD_FLAG) or "bold" in str(s["font"]).lower()
                            for s in spans
                        )
                        text = "".join(str(s["text"]) for s in line["spans"]).strip()
                        data.lines.append(
                            _Line(text, size, bool(bold), x0, y0, y1, int(block["number"]), number)
                        )
            except Exception as error:
                warnings.append(f"Page {number} could not be fully parsed: {type(error).__name__}.")
            pages.append(data)
        return pages

    @staticmethod
    def _drop_running_lines(pages: list[_PageData]) -> int:
        """Remove page numbers, repeated headers/footers and dot-leader TOC lines."""
        removed = 0
        zone_counter: Counter[str] = Counter()
        for page in pages:
            seen: set[str] = set()
            for line in page.lines:
                if _in_edge_zone(line, page.height):
                    seen.add(_normalise_key(line.text))
            zone_counter.update(seen)
        repeat_threshold = max(2, -(-len(pages) // 2))  # at least half of the pages, ceil
        for page in pages:
            for line in page.lines:
                key = _normalise_key(line.text)
                in_zone = _in_edge_zone(line, page.height)
                if (
                    _DOT_LEADER.search(line.text)
                    or (in_zone and _PAGE_NUMBER_LINE.match(line.text.strip()))
                    or (in_zone and len(pages) >= 2 and zone_counter[key] >= repeat_threshold)
                ):
                    line.drop = True
                    removed += 1
        return removed

    @staticmethod
    def _table_of_contents(pdf: Any) -> list[tuple[int, str, int]]:
        try:
            return [
                (int(lvl), str(title), int(page)) for lvl, title, page in pdf.get_toc(simple=True)
            ]
        except Exception:
            return []

    # -- headings ----------------------------------------------------------------------------
    def _mark_headings(self, pages: list[_PageData], toc: list[tuple[int, str, int]]) -> str:
        if toc and self._mark_from_toc(pages, toc):
            return "toc"
        if self._mark_from_fonts(pages):
            return "font"
        return "none"

    @staticmethod
    def _mark_from_toc(pages: list[_PageData], toc: list[tuple[int, str, int]]) -> bool:
        by_page = {page.number: page for page in pages}
        matched = 0
        for level, title, page_number in toc:
            wanted = _title_key(title)
            if not wanted:
                continue
            for number in (page_number, page_number + 1, page_number - 1):
                page = by_page.get(number)
                if page is None:
                    continue
                line = next(
                    (
                        item
                        for item in page.lines
                        if not item.drop
                        and item.heading_level is None
                        and _title_key(item.text) == wanted
                    ),
                    None,
                )
                if line is not None:
                    line.heading_level = max(1, level)
                    line.heading_source = "toc"
                    matched += 1
                    break
        return matched >= max(1, int(len(toc) * 0.5))

    @staticmethod
    def _mark_from_fonts(pages: list[_PageData]) -> bool:
        weights: Counter[float] = Counter()
        for page in pages:
            for line in page.lines:
                if not line.drop:
                    weights[round(line.size, 1)] += len(line.text)
        if not weights:
            return False
        body = weights.most_common(1)[0][0]
        candidates: list[_Line] = []
        for page in pages:
            block_sizes: dict[int, list[_Line]] = {}
            for line in page.lines:
                if not line.drop:
                    block_sizes.setdefault(line.block, []).append(line)
            for lines in block_sizes.values():
                for line in lines:
                    if is_bullet(line.text):
                        continue
                    bigger = line.size >= body + 1.0
                    bold_only = (
                        line.bold
                        and line.size >= body - 0.5
                        and (len(lines) == 1 or all(item.bold for item in lines))
                    )
                    _, title = split_number(line.text)
                    if (bigger or bold_only) and looks_like_heading_text(title):
                        candidates.append(line)
        if not candidates:
            return False
        sizes = sorted({round(c.size, 1) for c in candidates if c.size >= body + 1.0}, reverse=True)
        for line in candidates:
            number, _ = split_number(line.text)
            if number is not None:
                level = number_depth(number)
            elif round(line.size, 1) in sizes:
                level = sizes.index(round(line.size, 1)) + 1
            else:  # bold body-size line: deepest heading level
                level = len(sizes) + 1
            line.heading_level = min(level, 6)
            line.heading_source = "font"
        return True

    # -- blocks ------------------------------------------------------------------------------
    def _build_blocks(
        self, pdf: Any, pages: list[_PageData], warnings: list[str]
    ) -> tuple[list[ParsedBlock], list[int]]:
        blocks: list[ParsedBlock] = []
        ocr_pages: list[int] = []
        for data in pages:
            kept = [line for line in data.lines if not line.drop]
            chars = sum(len(line.text) for line in kept) + sum(
                len(cell) for _, rows, _ in data.tables for row in rows for cell in row
            )
            if chars < self._ocr_min_chars and (data.has_images or chars == 0):
                if (
                    self._ocr_provider is not None
                    and self._ocr_max_pages is not None
                    and len(ocr_pages) >= self._ocr_max_pages
                ):
                    warnings.append(
                        f"Page {data.number} was not OCR'd: the {self._ocr_max_pages}-page OCR "
                        "budget is exhausted."
                    )
                    continue
                recovered = self._ocr_page(pdf, data.number, warnings)
                if recovered:
                    blocks.extend(recovered)
                    ocr_pages.append(data.number)
                    continue
                if chars == 0:
                    warnings.append(f"Page {data.number} contains no extractable text.")
            blocks.extend(self._page_blocks(data, kept))
        return blocks, ocr_pages

    def _ocr_page(self, pdf: Any, page_number: int, warnings: list[str]) -> list[ParsedBlock]:
        if self._ocr_provider is None:
            return []
        try:
            pixmap = pdf[page_number - 1].get_pixmap(dpi=self._ocr_dpi, alpha=False)
            result = self._ocr_provider.extract(pixmap.tobytes("png"), page_number)
        except Exception as error:
            # IngestionError messages are written to be shown to users; others reveal only the type.
            reason = str(error) if isinstance(error, IngestionError) else type(error).__name__
            warnings.append(f"OCR failed on page {page_number}: {reason}")
            return []
        warnings.extend(f"OCR page {page_number}: {message}" for message in result.warnings)
        text = result.text.strip()
        if not text:
            warnings.append(f"OCR found no text on page {page_number}.")
            return []
        position = SourcePosition(page=page_number)
        return [
            ParsedBlock(text=part.strip(), position=position, ocr=True)
            for part in re.split(r"\n\s*\n", text)
            if part.strip()
        ]

    @staticmethod
    def _page_blocks(data: _PageData, lines: list[_Line]) -> list[ParsedBlock]:
        position = SourcePosition(page=data.number)
        items: list[tuple[float, int, ParsedBlock]] = []
        order = 0
        current: list[_Line] = []

        def flush() -> None:
            nonlocal order
            if not current:
                return
            text = _join_lines([line.text for line in current])
            if text:
                items.append((current[0].y0, order, ParsedBlock(text, position)))
                order += 1
            current.clear()

        for line in lines:
            if line.heading_level is not None:
                flush()
                items.append(
                    (
                        line.y0,
                        order,
                        ParsedBlock(
                            line.text,
                            position,
                            kind=BlockKind.HEADING,
                            heading_level=line.heading_level,
                            heading_source=line.heading_source,
                        ),
                    )
                )
                order += 1
                continue
            if current and (line.block != current[-1].block):
                flush()
            current.append(line)
        flush()
        for y0, rows, caption in data.tables:
            table = TableData(rows, position, caption)
            items.append(
                (
                    y0,
                    10_000 + len(items),
                    ParsedBlock(
                        render_table_markdown(rows), position, [table], kind=BlockKind.TABLE
                    ),
                )
            )
        items.sort(key=lambda item: (round(item[0], 1), item[1]))
        return [item[2] for item in items]

    @staticmethod
    def _metadata(pdf: Any, page_count: int) -> dict[str, Any]:
        raw = dict(pdf.metadata or {})
        keep = ("title", "author", "subject", "keywords", "creator", "producer")
        metadata: dict[str, Any] = {key: raw[key] for key in keep if raw.get(key)}
        for key, label in (("creationDate", "created"), ("modDate", "modified")):
            if raw.get(key):
                metadata[label] = _pdf_date(str(raw[key]))
        metadata["page_count"] = page_count
        return metadata


def _in_edge_zone(line: _Line, height: float) -> bool:
    return line.y1 <= height * _EDGE_ZONE + 20 or line.y0 >= height * (1 - _EDGE_ZONE) - 20


def _pdf_date(value: str) -> str:
    match = re.match(r"D:(\d{4})(\d{2})?(\d{2})?", value)
    if not match:
        return value
    return "-".join(part for part in match.groups() if part)


def _looks_like_numbered_heading(line: str) -> bool:
    number, title = split_number(line)
    return number is not None and looks_like_heading_text(title)


def _join_lines(lines: list[str]) -> str:
    """Re-flow hard-wrapped lines into paragraphs.

    List items and anything shaped like a numbered heading stay on their own line (the normalizer
    decides, with sequence validation, whether it really is a heading); the line after a possible
    heading starts a new paragraph. Everything else is wrapped prose and is joined, de-hyphenated.
    """
    out: list[str] = []
    previous_was_heading = False
    for line in lines:
        candidate = _looks_like_numbered_heading(line)
        if not out or is_list_item(line) or candidate or previous_was_heading:
            out.append(line)
        elif out[-1].endswith("-") and line[:1].islower():
            out[-1] = out[-1][:-1] + line
        else:
            out[-1] = f"{out[-1]} {line}"
        previous_was_heading = candidate
    return "\n".join(out)


# --------------------------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------------------------
class DOCXParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            document = DocxDocument(io.BytesIO(content))
        except Exception as error:
            raise IngestionError("DOCX parser could not open uploaded source.") from error
        blocks: list[ParsedBlock] = []
        saw_heading_style = False
        position = SourcePosition()
        pending_list: list[str] = []

        def flush_list() -> None:
            if pending_list:
                blocks.append(ParsedBlock("\n".join(pending_list), position))
                pending_list.clear()

        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = DocxParagraph(child, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                level = _word_heading_level(paragraph.style.name if paragraph.style else "")
                if level is not None:
                    flush_list()
                    saw_heading_style = True
                    blocks.append(
                        ParsedBlock(
                            text,
                            position,
                            kind=BlockKind.HEADING,
                            heading_level=level,
                            heading_source="style",
                        )
                    )
                elif _is_word_list(paragraph):
                    pending_list.append(text if is_list_item(text) else f"- {text}")
                else:
                    flush_list()
                    blocks.append(ParsedBlock(text, position))
            elif tag == "tbl":
                flush_list()
                table = DocxTable(child, document)
                rows = [_word_row(row) for row in table.rows]
                rows = [row for row in rows if any(row)]
                if rows:
                    blocks.append(
                        ParsedBlock(
                            render_table_markdown(rows),
                            position,
                            [TableData(rows, position)],
                            kind=BlockKind.TABLE,
                        )
                    )
        flush_list()
        properties = document.core_properties
        metadata: dict[str, Any] = {}
        for attr, label in (
            ("title", "title"),
            ("author", "author"),
            ("subject", "subject"),
            ("revision", "revision"),
            ("created", "created"),
            ("modified", "modified"),
        ):
            value = getattr(properties, attr, None)
            if value:
                metadata[label] = value.date().isoformat() if hasattr(value, "date") else str(value)
        return ParsedDocument(
            blocks=blocks,
            structure_source="style" if saw_heading_style else "none",
            metadata=metadata,
        )


def _word_heading_level(style_name: str) -> int | None:
    match = re.match(r"(?:Heading|Überschrift)\s+(\d)", style_name or "", re.IGNORECASE)
    if match:
        return int(match.group(1))
    if (style_name or "").lower() == "title":
        return 0  # the document title sits above "Heading 1"
    return None


def _is_word_list(paragraph: DocxParagraph) -> bool:
    name = (paragraph.style.name if paragraph.style else "") or ""
    if name.lower().startswith("list"):
        return True
    properties = paragraph._p.pPr
    return properties is not None and properties.numPr is not None


def _word_row(row: Any) -> list[str]:
    cells: list[str] = []
    previous = None
    for cell in row.cells:
        if cell._tc is previous:  # merged cells repeat the same underlying cell
            continue
        previous = cell._tc
        cells.append(cell.text.strip())
    return cells


# --------------------------------------------------------------------------------------------
# Spreadsheets and CSV
# --------------------------------------------------------------------------------------------
class XLSXParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise IngestionError("XLSX parser could not open uploaded source.") from error
        blocks: list[ParsedBlock] = []
        position = SourcePosition()
        for worksheet in workbook.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            blocks.append(
                ParsedBlock(
                    worksheet.title,
                    position,
                    kind=BlockKind.HEADING,
                    heading_level=1,
                    heading_source="style",
                )
            )
            if rows:
                blocks.append(
                    ParsedBlock(
                        render_table_markdown(rows),
                        position,
                        [TableData(rows, position)],
                        kind=BlockKind.TABLE,
                    )
                )
        return ParsedDocument(blocks=blocks, structure_source="style")


class CSVParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            decoded = content.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(decoded)))
        except (UnicodeDecodeError, csv.Error) as error:
            raise IngestionError("CSV parser could not read uploaded source.") from error
        position = SourcePosition()
        table = TableData(rows, position)
        return ParsedDocument(
            blocks=[
                ParsedBlock(render_table_markdown(rows), position, [table], kind=BlockKind.TABLE)
            ]
        )


# --------------------------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------------------------
_MD_HEADING = re.compile(r"^(?P<hash>#{1,6})\s+(?P<title>.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")
_MD_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


class MarkdownParser(DocumentParser):
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as error:
            raise IngestionError("Markdown source must be UTF-8.") from error
        blocks: list[ParsedBlock] = []
        saw_heading = False
        paragraph: list[str] = []
        items: list[str] = []
        table_rows: list[list[str]] = []
        fence: list[str] | None = None
        offset = 0
        start = 0

        def position(end: int) -> SourcePosition:
            return SourcePosition(start_offset=start, end_offset=end)

        def flush_paragraph(end: int) -> None:
            if paragraph:
                blocks.append(ParsedBlock(" ".join(paragraph), position(end)))
                paragraph.clear()

        def flush_items(end: int) -> None:
            if items:
                blocks.append(ParsedBlock("\n".join(items), position(end)))
                items.clear()

        def flush_table(end: int) -> None:
            if table_rows:
                table = TableData(list(table_rows), position(end))
                blocks.append(
                    ParsedBlock(
                        render_table_markdown(table.rows),
                        position(end),
                        [table],
                        kind=BlockKind.TABLE,
                    )
                )
                table_rows.clear()

        for raw in text.splitlines(keepends=True):
            line = raw.rstrip("\r\n")
            end = offset + len(raw)
            stripped = line.strip()
            if fence is not None:
                if _MD_FENCE.match(line):
                    blocks.append(ParsedBlock("\n".join(fence), position(end)))
                    fence = None
                else:
                    fence.append(line)
                offset = end
                continue
            if _MD_FENCE.match(line):
                flush_paragraph(offset)
                flush_items(offset)
                flush_table(offset)
                fence = []
                start = offset
                offset = end
                continue
            heading = _MD_HEADING.match(line)
            if heading:
                flush_paragraph(offset)
                flush_items(offset)
                flush_table(offset)
                start = offset
                saw_heading = True
                blocks.append(
                    ParsedBlock(
                        heading.group("title").strip(),
                        position(end),
                        kind=BlockKind.HEADING,
                        heading_level=len(heading.group("hash")),
                        heading_source="markdown",
                    )
                )
            elif stripped.startswith("|") and stripped.endswith("|"):
                flush_paragraph(offset)
                flush_items(offset)
                if not table_rows:
                    start = offset
                if not _MD_TABLE_SEPARATOR.match(stripped):
                    table_rows.append([cell.strip() for cell in stripped.strip("|").split("|")])
            elif not stripped:
                flush_paragraph(offset)
                flush_items(offset)
                flush_table(offset)
            elif LIST_MARKER.match(line):
                flush_paragraph(offset)
                flush_table(offset)
                if not items:
                    start = offset
                items.append(stripped)
            elif items and raw[:1] in (" ", "\t"):  # indented continuation of a list item
                items[-1] = f"{items[-1]} {stripped}"
            else:
                flush_items(offset)
                flush_table(offset)
                soft_wrap = paragraph and stripped[:1].islower()
                if paragraph and not soft_wrap:
                    flush_paragraph(offset)
                if not paragraph:
                    start = offset
                paragraph.append(stripped)
            offset = end
        flush_paragraph(offset)
        flush_items(offset)
        flush_table(offset)
        if fence:
            blocks.append(ParsedBlock("\n".join(fence), SourcePosition(offset, offset)))
        return ParsedDocument(blocks=blocks, structure_source="markdown" if saw_heading else "none")


class ParserFactory:
    """Maps validated document formats to deterministic parser implementations."""

    def __init__(
        self,
        ocr_provider: OCRProvider | None = None,
        ocr_min_chars: int = 30,
        ocr_dpi: int = 200,
        ocr_max_pages: int | None = None,
    ) -> None:
        self.ocr_provider = ocr_provider
        self._parsers: dict[DocumentFormat, DocumentParser] = {
            DocumentFormat.PDF: PDFParser(ocr_provider, ocr_min_chars, ocr_dpi, ocr_max_pages),
            DocumentFormat.DOCX: DOCXParser(),
            DocumentFormat.XLSX: XLSXParser(),
            DocumentFormat.CSV: CSVParser(),
            DocumentFormat.MARKDOWN: MarkdownParser(),
        }

    def parser_for(self, document_format: DocumentFormat) -> DocumentParser:
        return self._parsers[document_format]
